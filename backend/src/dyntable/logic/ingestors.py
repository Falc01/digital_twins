"""
Ingestores para a camada de dados (dyntable).

Foco: mínimo de hardcodes possível.
- Detecção configurável de colunas geográficas.
- Geo é OPCIONAL: se o arquivo de origem não contiver colunas de lat/lon,
  nenhuma coluna "lat"/"lon" é criada e nenhum valor falso é injetado.
- Reutiliza TableManager + DynTable (Pure Python Core).

Uso principal pelo backend (FastAPI) via IngestorFactory.
"""

import os
import re
import openpyxl
from typing import Optional

from .table_manager import TableManager
from ..data._core import DynType
from shared.config import get_geo_candidates, INGEST_INFER_TYPES


class BaseIngestor:
    def ingest(self, file_path: str, table_name: str, mgr: TableManager) -> None:
        raise NotImplementedError


class ExcelIngestor(BaseIngestor):
    def ingest(self, file_path: str, table_name: str, mgr: TableManager) -> None:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active

        # 1. Cabeçalhos limpos
        raw_headers = []
        for cell in ws[1]:
            val = str(cell.value).strip() if cell.value else f"col_{cell.column}"
            raw_headers.append(val)

        clean_headers = [re.sub(r'[^a-zA-Z0-9_]', '_', h).lower() for h in raw_headers]

        # 2. Obter/criar tabela
        t, created = mgr.get_or_create(table_name)
        existing_cols = set(t.column_names)

        ignored_cols = {"id", "created_at"}

        # 3. Adicionar colunas vindas do arquivo (STRING por padrão)
        for col in clean_headers:
            if col in ignored_cols:
                continue
            if col not in existing_cols:
                dtype = DynType.AUTO if INGEST_INFER_TYPES else DynType.STRING
                t.add_column(col, dtype, nullable=True)
                existing_cols.add(col)

        # 4. Detecção de geo (config-driven, sem hardcode de valores ou nomes fixos)
        lat_candidates, lon_candidates = get_geo_candidates()
        lat_col: Optional[str] = None
        lon_col: Optional[str] = None

        for col in clean_headers:
            if col in ignored_cols:
                continue
            if lat_col is None and col in lat_candidates:
                lat_col = col
            if lon_col is None and col in lon_candidates:
                lon_col = col

        # Criar colunas geo APENAS se foram detectadas no arquivo de origem
        if lat_col and lat_col not in existing_cols:
            t.add_column(lat_col, DynType.FLOAT, nullable=True)
            existing_cols.add(lat_col)
        if lon_col and lon_col not in existing_cols:
            t.add_column(lon_col, DynType.FLOAT, nullable=True)
            existing_cols.add(lon_col)

        # 5. Processar linhas
        for row in ws.iter_rows(min_row=2, values_only=True):
            if all(v is None for v in row):
                continue

            row_data: dict[str, any] = {}
            for i, val in enumerate(row):
                if i < len(clean_headers):
                    col = clean_headers[i]
                    if col in ignored_cols:
                        continue
                    if val is not None:
                        row_data[col] = val
                    else:
                        row_data[col] = None

            # 6. Parse geo (somente se as colunas geo existirem no arquivo)
            if lat_col and lat_col in row_data:
                try:
                    row_data[lat_col] = float(row_data[lat_col]) if row_data[lat_col] is not None else None
                except (ValueError, TypeError):
                    row_data[lat_col] = None

            if lon_col and lon_col in row_data:
                try:
                    row_data[lon_col] = float(row_data[lon_col]) if row_data[lon_col] is not None else None
                except (ValueError, TypeError):
                    row_data[lon_col] = None

            # Inserir (o DynTable cuida de inferência adicional em _apply_cell)
            try:
                t.new_row(**row_data)
            except Exception:
                # Ignora linhas problemáticas para robustez com dados reais
                continue

        mgr.save(t)


import csv

class CSVIngestor(BaseIngestor):
    def ingest(self, file_path: str, table_name: str, mgr: TableManager) -> None:
        # Tenta ler o arquivo CSV. Encoding UTF-8 com fallback para latin-1
        try:
            with open(file_path, mode='r', encoding='utf-8') as f:
                content = f.read()
        except UnicodeDecodeError:
            with open(file_path, mode='r', encoding='latin-1') as f:
                content = f.read()

        delimiter = ';' if ';' in content.split('\n')[0] else ','
        lines = content.splitlines()
        
        reader = csv.reader(lines, delimiter=delimiter)
        try:
            raw_headers = next(reader)
        except StopIteration:
            return

        clean_headers = [re.sub(r'[^a-zA-Z0-9_]', '_', h.strip()).lower() for h in raw_headers]

        t, created = mgr.get_or_create(table_name)
        existing_cols = set(t.column_names)

        ignored_cols = {"id", "created_at"}

        for col in clean_headers:
            if col in ignored_cols:
                continue
            if col not in existing_cols:
                dtype = DynType.AUTO if INGEST_INFER_TYPES else DynType.STRING
                t.add_column(col, dtype, nullable=True)
                existing_cols.add(col)

        lat_candidates, lon_candidates = get_geo_candidates()
        lat_col: Optional[str] = None
        lon_col: Optional[str] = None

        for col in clean_headers:
            if col in ignored_cols:
                continue
            if lat_col is None and col in lat_candidates:
                lat_col = col
            if lon_col is None and col in lon_candidates:
                lon_col = col

        if lat_col and lat_col not in existing_cols:
            t.add_column(lat_col, DynType.FLOAT, nullable=True)
            existing_cols.add(lat_col)
        if lon_col and lon_col not in existing_cols:
            t.add_column(lon_col, DynType.FLOAT, nullable=True)
            existing_cols.add(lon_col)

        for row in reader:
            if not row or all(v.strip() == "" for v in row):
                continue

            row_data: dict[str, any] = {}
            for i, val in enumerate(row):
                if i < len(clean_headers):
                    col = clean_headers[i]
                    if col in ignored_cols:
                        continue
                    val_str = val.strip()
                    if val_str == "":
                        row_data[col] = None
                    else:
                        row_data[col] = val_str

            if lat_col and lat_col in row_data:
                try:
                    row_data[lat_col] = float(row_data[lat_col]) if row_data[lat_col] is not None else None
                except (ValueError, TypeError):
                    row_data[lat_col] = None

            if lon_col and lon_col in row_data:
                try:
                    row_data[lon_col] = float(row_data[lon_col]) if row_data[lon_col] is not None else None
                except (ValueError, TypeError):
                    row_data[lon_col] = None

            try:
                t.new_row(**row_data)
            except Exception:
                continue

        mgr.save(t)


class IngestorFactory:
    @staticmethod
    def process_file(file_path: str, table_name: str, mgr: TableManager) -> None:
        ext = os.path.splitext(file_path)[1].lower()
        if ext in ('.xlsx', '.xls'):
            ingestor = ExcelIngestor()
        elif ext == '.csv':
            ingestor = CSVIngestor()
        else:
            raise ValueError(f"Formato {ext} não suportado pelo Ingestor.")

        ingestor.ingest(file_path, table_name, mgr)
