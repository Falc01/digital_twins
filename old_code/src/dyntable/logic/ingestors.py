import os
import openpyxl
import re
from src.dyntable.logic.table_manager import TableManager
from src.dyntable.data._core import DynType

class BaseIngestor:
    def ingest(self, file_path: str, table_name: str, mgr: TableManager) -> None:
        raise NotImplementedError

class ExcelIngestor(BaseIngestor):
    def ingest(self, file_path: str, table_name: str, mgr: TableManager) -> None:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active
        
        # Leitura da primeira linha para cabeçalhos
        headers = []
        for cell in ws[1]:
            val = str(cell.value).strip() if cell.value else f"col_{cell.column}"
            headers.append(val)
            
        clean_headers = [re.sub(r'[^a-zA-Z0-9_]', '_', h).lower() for h in headers]
        
        t, created = mgr.get_or_create(table_name)
        existing_cols = t.column_names
        
        for col in clean_headers:
            if col not in existing_cols:
                t.add_column(col, DynType.STRING, nullable=True)
                
        # Garantir colunas de coordenadas se não existirem
        if "lat" not in existing_cols:
            t.add_column("lat", DynType.FLOAT, nullable=True)
        if "lon" not in existing_cols:
            t.add_column("lon", DynType.FLOAT, nullable=True)

        POINTS = [
            (-12.9735, -38.5102),
            (-12.9725, -38.5120),
            (-12.9725, -38.5085),
            (-12.9745, -38.5120),
            (-12.9745, -38.5085),
        ]

        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
            # Se a linha for completamente vazia, pular
            if all(v is None for v in row):
                continue
                
            row_data = {}
            for i, val in enumerate(row):
                if i < len(clean_headers):
                    row_data[clean_headers[i]] = str(val) if val is not None else ""
            
            # Valores fixos de Pelourinho espalhados
            point = POINTS[idx % 5]
            row_data["lat"] = point[0]
            row_data["lon"] = point[1]
            
            t.new_row(**row_data)

        mgr.save(t)

class IngestorFactory:
    @staticmethod
    def process_file(file_path: str, table_name: str, mgr: TableManager):
        ext = os.path.splitext(file_path)[1].lower()
        if ext in ['.xlsx', '.xls']:
            ingestor = ExcelIngestor()
        else:
            raise ValueError(f"Formato {ext} não suportado pelo Ingestor.")
        
        ingestor.ingest(file_path, table_name, mgr)
